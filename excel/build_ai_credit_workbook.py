"""產生「美元 AI 相關公司債監控」Excel 活頁簿（Bloomberg Excel Add-in 版）。

執行（在 repo 根目錄）：
    python excel/build_ai_credit_workbook.py               # 產生 excel/AI_USD_Credit_Monitor.xlsx
    python excel/build_ai_credit_workbook.py --sample out.xlsx
        # 測試用：把 BDP/BDH 換成模擬數字，用來在沒有 Bloomberg 的環境驗證計算公式。
        # 模擬數字不是真實資料，不要拿來做任何判斷。

方法說明見 ../credit-quant.md。
"""

import argparse
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
FIRST, LAST = 5, 404  # 債券表資料列（400 檔）
ISS_FIRST, ISS_LAST = 5, 44  # 發行人表資料列（40 家，含空白可自行新增）
TOP_N = 20

BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT)
GREEN = Font(name=FONT, color="008000")
GREY = Font(name=FONT, color="808080")
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14)
HEAD = Font(name=FONT, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HELP_FILL = PatternFill("solid", fgColor="7F7F7F")
YELLOW = PatternFill("solid", fgColor="FFFF00")
LIGHT = PatternFill("solid", fgColor="F2F2F2")
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))
WRAP = Alignment(wrap_text=True, vertical="top")

# ---------------------------------------------------------------------------
# 發行人清單：分類、公司、股票代碼、債券 Ticker（用來對應債券的 TICKER 欄位）、備註
# 「請確認」的債券 Ticker：發債主體與母公司不同或代碼不確定，請在終端機用 DES 確認。
# ---------------------------------------------------------------------------
ISSUERS = [
    ("超大型雲端", "Microsoft", "MSFT US Equity", "MSFT", "AAA 級，發債量少"),
    ("超大型雲端", "Alphabet", "GOOGL US Equity", "GOOGL", "AI 資本支出擴大後恢復發債"),
    ("超大型雲端", "Amazon", "AMZN US Equity", "AMZN", ""),
    ("超大型雲端", "Meta Platforms", "META US Equity", "META", "大量發債支應資料中心；另有表外 SPV 融資"),
    ("超大型雲端", "Oracle", "ORCL US Equity", "ORCL", "槓桿高、AI 資本支出靠舉債，留意評等與 CDS"),
    ("大型科技", "Apple", "AAPL US Equity", "AAPL", ""),
    ("大型科技", "IBM", "IBM US Equity", "IBM", ""),
    ("半導體", "NVIDIA", "NVDA US Equity", "NVDA", "淨現金，流通債券少"),
    ("半導體", "Broadcom", "AVGO US Equity", "AVGO", "併購融資，發債量大"),
    ("半導體", "AMD", "AMD US Equity", "AMD", ""),
    ("半導體", "Intel", "INTC US Equity", "INTC", "評等下調壓力"),
    ("半導體", "Qualcomm", "QCOM US Equity", "QCOM", ""),
    ("半導體", "Micron", "MU US Equity", "MU", "HBM 記憶體，景氣循環"),
    ("半導體", "Marvell", "MRVL US Equity", "MRVL", ""),
    ("半導體", "Applied Materials", "AMAT US Equity", "AMAT", "半導體設備"),
    ("半導體", "TSMC", "2330 TT Equity", "請確認", "美元債由海外子公司發行，財報為台幣（已用美元覆寫）"),
    ("半導體", "SK hynix", "000660 KS Equity", "請確認", "財報為韓元（已用美元覆寫）"),
    ("伺服器／網通", "Cisco", "CSCO US Equity", "CSCO", ""),
    ("伺服器／網通", "Dell Technologies", "DELL US Equity", "DELL", "AI 伺服器"),
    ("伺服器／網通", "Hewlett Packard Enterprise", "HPE US Equity", "HPE", ""),
    ("軟體", "Salesforce", "CRM US Equity", "CRM", ""),
    ("資料中心", "Equinix", "EQIX US Equity", "EQIX", "REIT"),
    ("資料中心", "Digital Realty", "DLR US Equity", "DLR", "REIT"),
    ("AI 雲端（高收益）", "CoreWeave", "CRWV US Equity", "請確認", "高收益，債務以 GPU 擔保融資為主"),
]

RATINGS = [
    ("AAA", "AA以上"), ("AA+", "AA以上"), ("AA", "AA以上"), ("AA-", "AA以上"),
    ("A+", "A"), ("A", "A"), ("A-", "A"),
    ("BBB+", "BBB"), ("BBB", "BBB"), ("BBB-", "BBB"),
    ("BB+", "HY"), ("BB", "HY"), ("BB-", "HY"), ("B+", "HY"), ("B", "HY"), ("B-", "HY"),
    ("CCC+", "HY"), ("CCC", "HY"), ("CCC-", "HY"), ("CC", "HY"), ("C", "HY"), ("D", "HY"),
]
RATING_FIRST = 24
RATING_LAST = RATING_FIRST + len(RATINGS) - 1

P = "'參數'"
B = "'債券'"
I = "'發行人'"
S = "'分群統計'"
BUCKETS = ["AA以上", "A", "BBB", "HY"]
MAT_BUCKETS = ["1-3y", "3-5y", "5-7y", "7-10y", "10y+"]
# 分群統計表的位置
ST_R1, ST_R2 = 5, 5 + len(BUCKETS) - 1          # 評等分群統計
CM_HEAD = 12                                     # Carry 矩陣標題列
CM_R1, CM_R2 = 13, 13 + len(BUCKETS) - 1
IS_HEAD = 19                                     # 發行人曲線
IS_R1, IS_R2 = 20, 20 + (ISS_LAST - ISS_FIRST)

# 參數表的儲存格位置
ASOF, D1M, D3M, D6M = f"{P}!$B$4", f"{P}!$B$5", f"{P}!$B$6", f"{P}!$B$7"
MIN_GRP, MIN_ISS, MIN_RC = f"{P}!$B$9", f"{P}!$B$10", f"{P}!$B$11"
W_C, W_V, W_M, W_E = f"{P}!$B$13", f"{P}!$B$14", f"{P}!$B$15", f"{P}!$B$16"
TH_BUY, TH_SELL = f"{P}!$B$18", f"{P}!$B$19"
R_RTG = f"{P}!$A${RATING_FIRST}:$A${RATING_LAST}"
R_SCORE = f"{P}!$B${RATING_FIRST}:$B${RATING_LAST}"
R_BUCKET = f"{P}!$C${RATING_FIRST}:$C${RATING_LAST}"

# ---------------------------------------------------------------------------
# 債券表欄位：(key, 標題, 種類, 彭博欄位或說明)
# ---------------------------------------------------------------------------
BOND_COLS = [
    ("id", "債券代碼\n（ISIN 或彭博代碼）", "input", "← 貼上"),
    ("ticker", "發行人代碼", "bdp", "TICKER"),
    ("des", "債券簡稱", "bdp", "SECURITY_DES"),
    ("cpn", "票息 (%)", "bdp", "CPN"),
    ("mat", "到期日", "bdp", "MATURITY"),
    ("amt", "流通在外 (USD)", "bdp", "AMT_OUTSTANDING"),
    ("rtg", "彭博綜合評等", "bdp", "BB_COMPOSITE"),
    ("sp", "S&P", "bdp", "RTG_SP"),
    ("mdy", "Moody's", "bdp", "RTG_MOODY"),
    ("rank", "求償順位", "bdp", "PAYMENT_RANK"),
    ("mtytyp", "到期型態", "bdp", "MTY_TYP"),
    ("px", "價格 (Mid)", "bdp", "PX_MID"),
    ("ytm", "殖利率 (%)", "bdp", "YLD_YTM_MID"),
    ("oas", "OAS (bp)", "bdp", "OAS_SPREAD_MID"),
    ("dur", "利差存續期間", "bdp", "DUR_ADJ_MID"),
    ("oas1m", "OAS\n1 個月前", "bdh", D1M),
    ("oas3m", "OAS\n3 個月前", "bdh", D3M),
    ("oas6m", "OAS\n6 個月前", "bdh", D6M),
    ("cat", "分類", "calc", ""),
    ("eq6m", "股票 6M\n報酬 (%)", "calc", ""),
    ("yrs", "剩餘年限", "calc", ""),
    ("mb", "天期分群", "calc", ""),
    ("rscore", "評等分數", "calc", ""),
    ("rb", "評等分群", "calc", ""),
    ("valid", "有效\n(1/0)", "calc", ""),
    ("dts", "DTS", "calc", ""),
    ("d3m", "3M OAS\n變化 (bp)", "calc", ""),
    ("mom", "動能：6M→1M\n超額報酬 (bp)", "calc", ""),
    ("cn", "Carry\n組內檔數", "calc", ""),
    ("cmu", "Carry\n組平均 OAS", "calc", ""),
    ("csd", "Carry\n組標準差", "calc", ""),
    ("cz", "Carry z", "calc", ""),
    ("in", "發行人\n檔數", "calc", ""),
    ("ifit", "發行人曲線\n擬合 OAS", "calc", ""),
    ("ires", "相對發行人\n曲線 (bp)", "calc", ""),
    ("rfit", "評等曲線\n擬合 OAS", "calc", ""),
    ("rres", "相對評等\n曲線 (bp)", "calc", ""),
    ("vraw", "Value：\n相對便宜 %", "calc", ""),
    ("vz", "Value z", "calc", ""),
    ("mraw", "動能 ÷ DTS\n(×100)", "calc", ""),
    ("mz", "Momentum z", "calc", ""),
    ("ez", "股票動能 z", "calc", ""),
    ("score", "綜合分數", "calc", ""),
    ("pct", "評等分群內\n排名百分位", "calc", ""),
    ("signal", "訊號", "calc", ""),
    ("h_oas", "輔助\nOAS", "helper", ""),
    ("h_x", "輔助\nln(存續)", "helper", ""),
    ("h_amt", "輔助\n流通在外", "helper", ""),
    ("h_dts", "輔助\nDTS", "helper", ""),
    ("h_v", "輔助\nValue", "helper", ""),
    ("h_vok", "輔助\nValue 有效", "helper", ""),
    ("h_m", "輔助\n動能", "helper", ""),
    ("h_mok", "輔助\n動能有效", "helper", ""),
    ("h_e", "輔助\n股票", "helper", ""),
    ("h_eok", "輔助\n股票有效", "helper", ""),
    ("h_hi", "輔助\n排序（高）", "helper", ""),
    ("h_lo", "輔助\n排序（低）", "helper", ""),
]
C = {k: get_column_letter(i + 1) for i, (k, *_) in enumerate(BOND_COLS)}


def rng(key: str) -> str:
    """債券表某欄的資料範圍（絕對參照，給同一張表內的公式用）。"""
    return f"${C[key]}${FIRST}:${C[key]}${LAST}"


def xrng(key: str) -> str:
    """債券表某欄的資料範圍（給其他表引用）。"""
    return f"{B}!{rng(key)}"


def cell(key: str, r: int) -> str:
    return f"{C[key]}{r}"


def st(col: str, r1: int, r2: int) -> str:
    return f"{S}!${col}${r1}:${col}${r2}"


def zlookup(val: str, ok: str, r: int, n_col: str, mu_col: str, sd_col: str) -> str:
    """查「分群統計」表該評等分群的平均與標準差，計算 z 分數（截尾 ±3；分群少於 3 檔不計）。"""
    m = f"MATCH({cell('rb', r)},{st('A', ST_R1, ST_R2)},0)"
    n, mu, sd = (f"INDEX({st(x, ST_R1, ST_R2)},{m})" for x in (n_col, mu_col, sd_col))
    return (
        f'=IF({cell(ok, r)}<>1,"",IFERROR(IF({n}<3,"",'
        f"IF({sd}>0,MAX(-3,MIN(3,({cell(val, r)}-{mu})/{sd})),0)),\"\"))"
    )


def bond_formula(key: str, r: int) -> str | None:
    c = lambda k: cell(k, r)
    a = f"$A{r}"
    if key == "cat":
        return f'=IF({a}="","",IFERROR(INDEX({I}!$A${ISS_FIRST}:$A${ISS_LAST},MATCH({c("ticker")},{I}!$D${ISS_FIRST}:$D${ISS_LAST},0)),"未對應"))'
    if key == "eq6m":
        return f'=IF({a}="","",IFERROR(INDEX({I}!$G${ISS_FIRST}:$G${ISS_LAST},MATCH({c("ticker")},{I}!$D${ISS_FIRST}:$D${ISS_LAST},0))+0,""))'
    if key == "yrs":
        return f'=IF(ISNUMBER({c("mat")}),YEARFRAC({ASOF},{c("mat")}),"")'
    if key == "mb":
        y = c("yrs")
        return f'=IF({y}="","",IF({y}<3,"1-3y",IF({y}<5,"3-5y",IF({y}<7,"5-7y",IF({y}<10,"7-10y","10y+")))))'
    if key == "rscore":
        return f'=IF({a}="","",IFERROR(INDEX({R_SCORE},MATCH({c("rtg")},{R_RTG},0)),""))'
    if key == "rb":
        return f'=IF({a}="","",IFERROR(INDEX({R_BUCKET},MATCH({c("rtg")},{R_RTG},0)),""))'
    if key == "valid":
        return (
            f'=IF(AND(ISNUMBER({c("oas")}),ISNUMBER({c("dur")}),ISNUMBER({c("yrs")})),'
            f'IF(AND({c("dur")}>0,{c("rb")}<>""),1,0),0)'
        )
    if key == "dts":
        return f'=IF({c("valid")}=1,{c("dur")}*{c("oas")},"")'
    if key == "d3m":
        return f'=IF({c("valid")}=1,IF(ISNUMBER({c("oas3m")}),{c("oas")}-{c("oas3m")},""),"")'
    if key == "mom":
        return (
            f'=IF({c("valid")}=1,IF(AND(ISNUMBER({c("oas1m")}),ISNUMBER({c("oas6m")})),'
            f'{c("oas6m")}*5/12-{c("dur")}*({c("oas1m")}-{c("oas6m")}),""),"")'
        )
    if key == "cn":
        return f'=IF({c("valid")}=1,COUNTIFS({rng("rb")},{c("rb")},{rng("mb")},{c("mb")},{rng("valid")},1),"")'
    # Carry：評等 × 天期分群的平均／標準差；組內檔數不足時退回評等分群
    rb_m = f"MATCH({c('rb')},{st('A', CM_R1, CM_R2)},0)"
    mb_m = lambda col1: f"MATCH({c('mb')},{S}!${col1}${CM_HEAD}:${chr(ord(col1) + 4)}${CM_HEAD},0)"
    mat_rng = lambda col1: f"{S}!${col1}${CM_R1}:${chr(ord(col1) + 4)}${CM_R2}"
    rb_s = f"MATCH({c('rb')},{st('A', ST_R1, ST_R2)},0)"
    if key == "cmu":
        return (
            f'=IF({c("valid")}=1,IFERROR(IF({c("cn")}>={MIN_GRP},INDEX({mat_rng("G")},{rb_m},{mb_m("G")}),'
            f'INDEX({st("I", ST_R1, ST_R2)},{rb_s})),""),"")'
        )
    if key == "csd":
        return (
            f'=IF({c("valid")}=1,IFERROR(IF({c("cn")}>={MIN_GRP},INDEX({mat_rng("L")},{rb_m},{mb_m("L")}),'
            f'INDEX({st("J", ST_R1, ST_R2)},{rb_s})),""),"")'
        )
    if key == "cz":
        return (
            f'=IF(ISNUMBER({c("csd")}),IF({c("csd")}>0,'
            f'MAX(-3,MIN(3,({c("oas")}-{c("cmu")})/{c("csd")})),0),"")'
        )
    if key == "in":
        return f'=IF({c("valid")}=1,COUNTIFS({rng("ticker")},{c("ticker")},{rng("valid")},1),"")'
    if key == "ifit":
        m = f"MATCH({c('ticker')},{st('A', IS_R1, IS_R2)},0)"
        b_, a_ = f"INDEX({st('G', IS_R1, IS_R2)},{m})", f"INDEX({st('H', IS_R1, IS_R2)},{m})"
        return f'=IF({c("valid")}=1,IFERROR(IF(ISNUMBER({b_}),{a_}+{b_}*{c("h_x")},""),""),"")'
    if key == "ires":
        return f'=IF(ISNUMBER({c("ifit")}),{c("oas")}-{c("ifit")},"")'
    if key == "rfit":
        b_, a_ = f"INDEX({st('G', ST_R1, ST_R2)},{rb_s})", f"INDEX({st('H', ST_R1, ST_R2)},{rb_s})"
        return f'=IF({c("valid")}=1,IFERROR(IF(ISNUMBER({b_}),{a_}+{b_}*{c("h_x")},""),""),"")'
    if key == "rres":
        return f'=IF(ISNUMBER({c("rfit")}),{c("oas")}-{c("rfit")},"")'
    if key == "vraw":
        return f'=IF(ISNUMBER({c("rfit")}),IF({c("rfit")}>0,{c("rres")}/{c("rfit")},""),"")'
    if key == "vz":
        return zlookup("h_v", "h_vok", r, "K", "L", "M")
    if key == "mz":
        return zlookup("h_m", "h_mok", r, "N", "O", "P")
    if key == "ez":
        return zlookup("h_e", "h_eok", r, "Q", "R", "S")
    if key == "mraw":
        return f'=IF(ISNUMBER({c("mom")}),IF({c("dts")}>0,{c("mom")}/{c("dts")}*100,""),"")'
    if key == "mz":
        return zscore("h_m", "h_mok", r)
    if key == "ez":
        return zscore("h_e", "h_eok", r)
    if key == "score":
        z = lambda k: f"IF(ISNUMBER({c(k)}),{c(k)},0)"
        return (
            f'=IF({c("valid")}=1,({z("cz")}*{W_C}+{z("vz")}*{W_V}+{z("mz")}*{W_M}+{z("ez")}*{W_E})'
            f'/({W_C}+{W_V}+{W_M}+{W_E}),"")'
        )
    if key == "pct":
        # 用數值比較（排序輔助欄含列號，避免同分），不用 COUNTIFS 的 ">"&分數：會把分數轉成文字而失去精度
        return (
            f'=IF({c("valid")}=1,(SUMPRODUCT(({rng("rb")}={c("rb")})*({rng("valid")}=1)*({rng("h_hi")}>{c("h_hi")}))+1)'
            f'/COUNTIFS({rng("rb")},{c("rb")},{rng("valid")},1),"")'
        )
    if key == "signal":
        p = c("pct")
        return f'=IF(ISNUMBER({p}),IF({p}<={TH_BUY},"加碼候選",IF({p}>{TH_SELL},"減碼候選","")),"")'
    if key == "h_oas":
        return f'=IF({c("valid")}=1,{c("oas")},0)'
    if key == "h_x":
        return f'=IF({c("valid")}=1,LN({c("dur")}),0)'
    if key == "h_amt":
        return f'=IF({c("valid")}=1,IF(ISNUMBER({c("amt")}),{c("amt")},0),0)'
    if key == "h_dts":
        return f'=IF({c("valid")}=1,{c("dts")},0)'
    for h, src in (("v", "vraw"), ("m", "mraw"), ("e", "eq6m")):
        if key == f"h_{h}":
            return f'=IF({c("valid")}=1,IF(ISNUMBER({c(src)}),{c(src)},0),0)'
        if key == f"h_{h}ok":
            return f'=IF({c("valid")}=1,IF(ISNUMBER({c(src)}),1,0),0)'
    if key == "h_hi":
        return f'=IF({c("valid")}=1,{c("score")}+ROW()/1E9,-1E9)'
    if key == "h_lo":
        return f'=IF({c("valid")}=1,{c("score")}+ROW()/1E9,1E9)'
    return None


# ---------------------------------------------------------------------------
# 模擬資料（只在 --sample 模式使用）
# ---------------------------------------------------------------------------
def sample_bond_values(n_rows: int, seed: int = 1) -> list[dict]:
    rng_ = np.random.default_rng(seed)
    rtg_by_iss = {"MSFT": "AAA", "GOOGL": "AA+", "AMZN": "AA", "META": "AA-", "ORCL": "BBB", "AAPL": "AA+",
                  "IBM": "A-", "NVDA": "AA-", "AVGO": "BBB+", "AMD": "A", "INTC": "BBB", "QCOM": "A",
                  "MU": "BBB-", "MRVL": "BBB-", "AMAT": "A", "CSCO": "AA-", "DELL": "BBB", "HPE": "BBB",
                  "CRM": "A+", "EQIX": "BBB+", "DLR": "BBB"}
    base = {"AA以上": 45, "A": 70, "BBB": 105}
    bucket = dict(RATINGS)
    out = []
    for i in range(n_rows):
        t = rng_.choice(list(rtg_by_iss))
        rtg = rtg_by_iss[t]
        yrs = float(rng_.uniform(1.2, 30))
        dur = min(yrs, 22) * 0.82
        oas = base[bucket[rtg]] * (1 + 0.1 * np.sqrt(yrs)) * np.exp(rng_.normal(0, 0.12))
        out.append({
            "id": f"TEST{i:04d} Corp", "ticker": t, "des": f"{t} {rng_.uniform(1, 6):.2f} 20{25 + int(yrs):02d}",
            "cpn": round(float(rng_.uniform(1, 6)), 3), "mat_years": yrs,
            "amt": float(rng_.choice([5e8, 7.5e8, 1e9, 1.5e9, 2e9])), "rtg": rtg, "sp": rtg, "mdy": "",
            "rank": "Sr Unsecured", "mtytyp": "CALLABLE", "px": round(float(rng_.uniform(70, 105)), 3),
            "ytm": round(4 + oas / 100, 3), "oas": round(oas, 1), "dur": round(dur, 2),
            "oas1m": round(oas * np.exp(rng_.normal(0, 0.05)), 1),
            "oas3m": round(oas * np.exp(rng_.normal(0, 0.08)), 1),
            "oas6m": round(oas * np.exp(rng_.normal(0, 0.1)), 1),
        })
    out[3]["id"] = ""  # 模擬中間有空白列
    out[7]["rtg"] = "NR"  # 模擬無評等
    return out


# ---------------------------------------------------------------------------
def style_header(ws, row, c1, c2, fill=HEAD_FILL):
    for col in range(c1, c2 + 1):
        x = ws.cell(row=row, column=col)
        x.font, x.fill = HEAD, fill
        x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def build_params(wb):
    ws = wb.create_sheet("參數")
    ws["A1"], ws["A1"].font = "參數設定", TITLE
    ws["A2"], ws["A2"].font = "黃底藍字為可調整的參數。", GREY
    rows = [
        (4, "評價日", "=TODAY()", "yyyy-mm-dd", "預設為今天；要回看特定日期可改成固定日期"),
        (5, "1 個月前", "=EDATE(B4,-1)", "yyyy-mm-dd", "動能與 OAS 變化的比較日期"),
        (6, "3 個月前", "=EDATE(B4,-3)", "yyyy-mm-dd", ""),
        (7, "6 個月前", "=EDATE(B4,-6)", "yyyy-mm-dd", ""),
        (9, "Carry 分群最少檔數", 5, "0", "「評等 × 天期」分群不足此檔數時，改用評等分群"),
        (10, "發行人曲線最少檔數", 3, "0", "同一發行人至少幾檔才擬合發行人曲線"),
        (11, "評等曲線最少檔數", 5, "0", "同一評等分群至少幾檔才擬合評等曲線"),
        (13, "權重：Carry", 1, "0.00", "綜合分數 = 各 z 分數依權重加權平均"),
        (14, "權重：Value", 1, "0.00", ""),
        (15, "權重：Momentum（債券）", 0.5, "0.00", ""),
        (16, "權重：股票動能", 0.5, "0.00", "投資等級債的動能以股票訊號為主"),
        (18, "加碼候選門檻（百分位 ≤）", 0.2, "0%", "評等分群內綜合分數前 20%"),
        (19, "減碼候選門檻（百分位 >）", 0.8, "0%", "評等分群內綜合分數後 20%"),
    ]
    for r, label, val, fmt, note in rows:
        ws.cell(row=r, column=1, value=label).font = BLACK
        x = ws.cell(row=r, column=2, value=val)
        x.font, x.fill, x.number_format = BLUE, YELLOW, fmt
        ws.cell(row=r, column=3, value=note).font = GREY
    ws.cell(row=RATING_FIRST - 2, column=1, value="評等對照表").font = BOLD
    for i, h in enumerate(["評等（BB_COMPOSITE）", "評等分數", "評等分群"]):
        ws.cell(row=RATING_FIRST - 1, column=i + 1, value=h)
    style_header(ws, RATING_FIRST - 1, 1, 3)
    for i, (rtg, bkt) in enumerate(RATINGS):
        r = RATING_FIRST + i
        ws.cell(row=r, column=1, value=rtg).font = BLUE
        ws.cell(row=r, column=2, value=i + 1).font = BLUE
        ws.cell(row=r, column=3, value=bkt).font = BLUE
    ws.cell(row=RATING_LAST + 1, column=1,
            value="評等分數：1 = AAA，數字越大評等越低。對照表以外的評等（例如 NR）不列入計算。").font = GREY
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 60


def build_issuers(wb, sample: bool):
    ws = wb.create_sheet("發行人")
    ws["A1"], ws["A1"].font = "發行人：AI 相關主流公司", TITLE
    ws["A2"] = ("A–E 欄可自行增刪（黃底）。D 欄要和債券的 TICKER 一致才能對應；標「請確認」的請在終端機用 DES 查發債主體代碼。"
                "財報欄位以美元覆寫 (EQY_FUND_CRNCY=USD)，單位為百萬美元。")
    ws["A2"].font = GREY
    cols = [
        ("分類", "input", None, 18), ("公司", "input", None, 24), ("股票代碼", "input", None, 18),
        ("債券 Ticker", "input", None, 12), ("備註", "input", None, 36),
        ("市值\n(百萬 USD)", "eq", "CUR_MKT_CAP", 14), ("股價 6M\n報酬 (%)", "eq", "CHG_PCT_6M", 11),
        ("90 天\n波動度 (%)", "eq", "VOLATILITY_90D", 11), ("營收", "fund", "SALES_REV_TURN", 12),
        ("資本支出", "fund", "CAPITAL_EXPEND", 12), ("自由現金流", "fund", "CF_FREE_CASH_FLOW", 12),
        ("總負債", "fund", "SHORT_AND_LONG_TERM_DEBT", 12), ("淨負債", "fund", "NET_DEBT", 12),
        ("EBITDA", "fund", "EBITDA", 12), ("CDS 5Y 代碼", "eq", "CDS_SPREAD_TICKER_5Y", 18),
        ("CDS 5Y (bp)", "cds", "PX_LAST", 11),
        ("淨負債 /\nEBITDA (x)", "calc", None, 11), ("資本支出 /\n營收", "calc", None, 11),
        ("自由現金流 /\n總負債", "calc", None, 12), ("債券檔數", "calc", None, 9),
        ("流通在外合計\n(百萬 USD)", "calc", None, 14), ("加權平均\nOAS (bp)", "calc", None, 11),
        ("加權平均\nDTS", "calc", None, 10), ("平均\n綜合分數", "calc", None, 10),
        ("加碼候選\n檔數", "calc", None, 10), ("減碼候選\n檔數", "calc", None, 10),
    ]
    ws.cell(row=3, column=1, value="彭博欄位 →").font = GREY
    for j, (h, kind, fld, w) in enumerate(cols, start=1):
        ws.cell(row=4, column=j, value=h)
        ws.column_dimensions[get_column_letter(j)].width = w
        if fld:
            x = ws.cell(row=3, column=j, value=fld)
            x.font, x.fill = BLUE, YELLOW
    style_header(ws, 4, 1, len(cols))
    ws.row_dimensions[4].height = 42

    srng = np.random.default_rng(7)
    for i in range(ISS_LAST - ISS_FIRST + 1):
        r = ISS_FIRST + i
        if i < len(ISSUERS):
            for j, v in enumerate(ISSUERS[i], start=1):
                x = ws.cell(row=r, column=j, value=v)
                x.font = BLUE
                if v == "請確認":
                    x.fill = YELLOW
        for j, (h, kind, fld, w) in enumerate(cols, start=1):
            L = get_column_letter(j)
            x = ws.cell(row=r, column=j)
            if kind in ("eq", "fund", "cds"):
                if sample:
                    x.value = None if i >= len(ISSUERS) else {
                        "CUR_MKT_CAP": float(srng.uniform(5e4, 4e6)), "CHG_PCT_6M": float(srng.normal(10, 25)),
                        "VOLATILITY_90D": float(srng.uniform(20, 60)), "SALES_REV_TURN": float(srng.uniform(1e4, 6e5)),
                        "CAPITAL_EXPEND": -float(srng.uniform(1e3, 8e4)), "CF_FREE_CASH_FLOW": float(srng.normal(2e4, 2e4)),
                        "SHORT_AND_LONG_TERM_DEBT": float(srng.uniform(1e3, 1.2e5)), "NET_DEBT": float(srng.normal(1e4, 4e4)),
                        "EBITDA": float(srng.uniform(2e3, 2e5)), "CDS_SPREAD_TICKER_5Y": "TEST CDS",
                        "PX_LAST": float(srng.uniform(20, 150)),
                    }[fld]
                elif kind == "eq":
                    x.value = f'=IF($C{r}="","",BDP($C{r},{L}$3))'
                elif kind == "fund":
                    x.value = f'=IF($C{r}="","",BDP($C{r},{L}$3,"EQY_FUND_CRNCY","USD"))'
                else:
                    x.value = f'=IF(LEFT($O{r},1)="#","",IF($O{r}="","",BDP($O{r},{L}$3)))'
                x.font = BLACK
            elif kind == "calc":
                tk = f"$D{r}"
                m = f"({xrng('ticker')}={tk})*({xrng('valid')}=1)"
                f = {
                    "淨負債 /\nEBITDA (x)": f'=IF(AND(ISNUMBER($M{r}),ISNUMBER($N{r})),IF($N{r}>0,$M{r}/$N{r},""),"")',
                    "資本支出 /\n營收": f'=IF(AND(ISNUMBER($J{r}),ISNUMBER($I{r})),IF($I{r}>0,ABS($J{r})/$I{r},""),"")',
                    "自由現金流 /\n總負債": f'=IF(AND(ISNUMBER($K{r}),ISNUMBER($L{r})),IF($L{r}>0,$K{r}/$L{r},""),"")',
                    "債券檔數": f'=IF({tk}="","",COUNTIFS({xrng("ticker")},{tk},{xrng("valid")},1))',
                    "流通在外合計\n(百萬 USD)": f'=IF({tk}="","",SUMPRODUCT({m}*{xrng("h_amt")})/1E6)',
                    "加權平均\nOAS (bp)": f'=IF({tk}="","",IF(SUMPRODUCT({m}*{xrng("h_amt")})>0,SUMPRODUCT({m}*{xrng("h_amt")}*{xrng("h_oas")})/SUMPRODUCT({m}*{xrng("h_amt")}),""))',
                    "加權平均\nDTS": f'=IF({tk}="","",IF(SUMPRODUCT({m}*{xrng("h_amt")})>0,SUMPRODUCT({m}*{xrng("h_amt")}*{xrng("h_dts")})/SUMPRODUCT({m}*{xrng("h_amt")}),""))',
                    "平均\n綜合分數": f'=IF({tk}="","",IFERROR(AVERAGEIFS({xrng("score")},{xrng("ticker")},{tk},{xrng("valid")},1),""))',
                    "加碼候選\n檔數": f'=IF({tk}="","",COUNTIFS({xrng("ticker")},{tk},{xrng("signal")},"加碼候選"))',
                    "減碼候選\n檔數": f'=IF({tk}="","",COUNTIFS({xrng("ticker")},{tk},{xrng("signal")},"減碼候選"))',
                }[h]
                x.value = f
                x.font = GREEN if B in f else BLACK
            else:
                x.fill = YELLOW
            fmt = {"市值\n(百萬 USD)": "#,##0", "營收": "#,##0", "資本支出": "#,##0;(#,##0);-",
                   "自由現金流": "#,##0;(#,##0);-", "總負債": "#,##0", "淨負債": "#,##0;(#,##0);-",
                   "EBITDA": "#,##0;(#,##0);-", "流通在外合計\n(百萬 USD)": "#,##0",
                   "淨負債 /\nEBITDA (x)": "0.0x;(0.0x);-", "資本支出 /\n營收": "0.0%", "自由現金流 /\n總負債": "0.0%"}
            x.number_format = fmt.get(h, "0.0" if kind in ("eq", "cds", "calc") else "General")
            if h in ("債券檔數", "加碼候選\n檔數", "減碼候選\n檔數"):
                x.number_format = "0"
            if h == "平均\n綜合分數":
                x.number_format = "0.00"
    ws.freeze_panes = "C5"
    last = get_column_letter(len(cols))
    ws.conditional_formatting.add(
        f"X{ISS_FIRST}:X{ISS_LAST}",
        ColorScaleRule(start_type="num", start_value=-1, start_color="F8696B",
                       mid_type="num", mid_value=0, mid_color="FFFFFF",
                       end_type="num", end_value=1, end_color="63BE7B"))
    ws.auto_filter.ref = f"A4:{last}{ISS_LAST}"


def build_bonds(wb, sample: bool):
    ws = wb.create_sheet("債券")
    ws["A1"], ws["A1"].font = "美元公司債：AI 相關發行人", TITLE
    ws["A2"] = ("A 欄貼上債券代碼（見「說明」）。第 3 列黃底是彭博欄位代碼，可依 FLDS 查詢結果修改。"
                "灰色標題為輔助欄，請勿刪除。")
    ws["A2"].font = GREY
    ws.cell(row=3, column=1, value="彭博欄位 →").font = GREY
    for j, (k, h, kind, fld) in enumerate(BOND_COLS, start=1):
        ws.cell(row=4, column=j, value=h)
        if kind == "bdp":
            x = ws.cell(row=3, column=j, value=fld)
            x.font, x.fill = BLUE, YELLOW
        elif kind == "bdh":
            x = ws.cell(row=3, column=j, value="BDH（同 OAS 欄位）")
            x.font = GREY
    style_header(ws, 4, 1, len(BOND_COLS))
    for j, (k, *_ ) in enumerate(BOND_COLS, start=1):
        if k.startswith("h_"):
            ws.cell(row=4, column=j).fill = HELP_FILL
    ws.row_dimensions[4].height = 44

    samples = sample_bond_values(120) if sample else None
    oas_col = C["oas"]
    for r in range(FIRST, LAST + 1):
        s = samples[r - FIRST] if sample and r - FIRST < len(samples) else None
        for j, (k, h, kind, fld) in enumerate(BOND_COLS, start=1):
            x = ws.cell(row=r, column=j)
            if kind == "input":
                x.font, x.fill = BLUE, YELLOW
                if s:
                    x.value = s["id"] or None
            elif kind in ("bdp", "bdh"):
                x.font = BLACK
                if sample:
                    if s and s["id"]:
                        x.value = f'=EDATE({ASOF},{round(s["mat_years"] * 12)})' if k == "mat" else s[k]
                elif kind == "bdp":
                    x.value = f'=IF($A{r}="","",BDP($A{r},{get_column_letter(j)}$3))'
                else:
                    x.value = f'=IF($A{r}="","",BDH($A{r},${oas_col}$3,{fld},{fld},"Days=A","Fill=P","Dates=H"))'
            else:
                x.value = bond_formula(k, r)
                x.font = GREY if kind == "helper" else (GREEN if I in (x.value or "") else BLACK)
            fmt = {"cpn": "0.000", "mat": "yyyy-mm-dd", "amt": "#,##0", "px": "0.000", "ytm": "0.000",
                   "oas": "0", "dur": "0.00", "oas1m": "0", "oas3m": "0", "oas6m": "0", "eq6m": "0.0",
                   "yrs": "0.0", "rscore": "0", "valid": "0", "dts": "#,##0", "d3m": "0;-0;0", "mom": "0;-0;0",
                   "cn": "0", "cmu": "0", "csd": "0.0", "in": "0", "ifit": "0", "ires": "0.0;-0.0;0.0",
                   "rfit": "0", "rres": "0.0;-0.0;0.0", "vraw": "0.0%", "mraw": "0.00", "pct": "0%",
                   "h_amt": "#,##0"}.get(k)
            if k in ("cz", "vz", "mz", "ez", "score"):
                fmt = "0.00"
            if fmt:
                x.number_format = fmt
    widths = {"id": 22, "ticker": 9, "des": 22, "rank": 14, "mtytyp": 12, "cat": 16, "signal": 11, "amt": 15,
              "mat": 11}
    for j, (k, *_ ) in enumerate(BOND_COLS, start=1):
        ws.column_dimensions[get_column_letter(j)].width = widths.get(k, 10)
    ws.freeze_panes = "D5"
    ws.auto_filter.ref = f"A4:{C['signal']}{LAST}"
    ws.conditional_formatting.add(
        f"{C['score']}{FIRST}:{C['score']}{LAST}",
        ColorScaleRule(start_type="num", start_value=-1.5, start_color="F8696B",
                       mid_type="num", mid_value=0, mid_color="FFFFFF",
                       end_type="num", end_value=1.5, end_color="63BE7B"))
    sig = f"{C['signal']}{FIRST}:{C['signal']}{LAST}"
    ws.conditional_formatting.add(sig, CellIsRule(operator="equal", formula=['"加碼候選"'],
                                                  fill=PatternFill("solid", fgColor="C6EFCE")))
    ws.conditional_formatting.add(sig, CellIsRule(operator="equal", formula=['"減碼候選"'],
                                                  fill=PatternFill("solid", fgColor="FFC7CE")))
    notes = {
        "dur": "以修正存續期間近似利差存續期間；可轉換或可贖回債請改用 OAS 存續期間欄位（以 FLDS 搜尋 OAS duration）。",
        "oas1m": "用 BDH 取該日的 OAS（非交易日以前一個交易日補值）。400 檔 × 3 個日期會用掉不少資料額度，建議每月更新一次。",
        "mom": "近似過去 6 個月到 1 個月前的超額報酬（bp）：5 個月的利差 carry − 利差存續期間 × OAS 變化。跳過最近 1 個月以避開短期反轉。",
        "cz": "評等 × 天期分群內，OAS 相對組平均的 z 分數（截尾 ±3）。高 = 利差較高。",
        "ires": "以同一發行人的債券擬合 OAS = a + b·ln(存續期間)，正值 = 相對自家曲線便宜。用於同一發行人內選券。",
        "rres": "以同評等分群的所有債券擬合 OAS = a + b·ln(存續期間)，正值 = 相對同評等曲線便宜。",
        "vraw": "相對評等曲線殘差 ÷ 擬合 OAS，例如 +10% 代表利差比同評等、同存續期間的公允水準高 10%。",
        "score": "Carry、Value、Momentum、股票動能 z 分數的加權平均，權重在「參數」表。高 = 偏便宜／偏強勢。",
        "pct": "在同評等分群內，綜合分數的排名百分位。0% = 最高分。",
    }
    for k, t in notes.items():
        ws[f"{C[k]}4"].comment = Comment(t, "Claude")


def build_stats(wb):
    """分群統計：每個評等分群、Carry 分群、發行人只計算一次，債券表再查表引用。"""
    ws = wb.create_sheet("分群統計")
    ws["A1"], ws["A1"].font = "分群統計（自動計算，請勿修改）", TITLE
    ws["A2"] = ("曲線擬合：OAS = 截距 + 斜率 × ln(存續期間)，用最小平方法；檔數不足「參數」表門檻時不擬合。"
                "z 分數用的平均與標準差也在這裡計算。")
    ws["A2"].font = GREY

    def put(r, col, f, fmt="0.00"):
        x = ws[f"{col}{r}"]
        x.value, x.number_format = f, fmt
        x.font = GREEN if B in str(f) or I in str(f) else BLACK

    def fit_cols(r, mask, min_n):
        sp = lambda e: f"=SUMPRODUCT({mask}{e})"
        put(r, "B", sp(""), "0")
        put(r, "C", sp(f"*{xrng('h_x')}"))
        put(r, "D", sp(f"*{xrng('h_oas')}"), "#,##0")
        put(r, "E", sp(f"*{xrng('h_x')}^2"))
        put(r, "F", sp(f"*{xrng('h_x')}*{xrng('h_oas')}"), "#,##0")
        den = f"(B{r}*E{r}-C{r}^2)"
        put(r, "G", f'=IF(B{r}<{min_n},"",IF(ABS({den})<1E-9,0,(B{r}*F{r}-C{r}*D{r})/{den}))', "0.0")
        put(r, "H", f'=IF(B{r}<{min_n},"",(D{r}-G{r}*C{r})/B{r})', "0.0")

    # A. 評等分群
    ws.cell(row=3, column=1, value="A. 評等分群：評等曲線、Carry 退回值、各因子的平均與標準差").font = BOLD
    heads = ["評等分群", "檔數", "Σx", "Σy", "Σx²", "Σxy", "曲線斜率", "曲線截距", "OAS 平均", "OAS 標準差",
             "Value 檔數", "Value 平均", "Value 標準差", "動能 檔數", "動能 平均", "動能 標準差",
             "股票 檔數", "股票 平均", "股票 標準差"]
    for j, h in enumerate(heads, start=1):
        ws.cell(row=4, column=j, value=h)
    style_header(ws, 4, 1, len(heads))
    for i, bkt in enumerate(BUCKETS):
        r = ST_R1 + i
        ws[f"A{r}"], ws[f"A{r}"].font = bkt, BLUE
        mask = f"({xrng('valid')}=1)*({xrng('rb')}=$A{r})"
        fit_cols(r, mask, MIN_RC)
        put(r, "I", f'=IF(B{r}>0,D{r}/B{r},"")', "0.0")
        put(r, "J", f'=IF(B{r}>1,SQRT(SUMPRODUCT({mask}*({xrng("h_oas")}-I{r})^2)/(B{r}-1)),"")', "0.0")
        for cn, cmu, csd, val, ok, fmt in (("K", "L", "M", "h_v", "h_vok", "0.0%"),
                                          ("N", "O", "P", "h_m", "h_mok", "0.00"),
                                          ("Q", "R", "S", "h_e", "h_eok", "0.0")):
            m = f"({xrng(ok)}=1)*({xrng('rb')}=$A{r})"
            put(r, cn, f"=SUMPRODUCT({m})", "0")
            put(r, cmu, f'=IF({cn}{r}>0,SUMPRODUCT({m}*{xrng(val)})/{cn}{r},"")', fmt)
            put(r, csd, f'=IF({cn}{r}>1,SQRT(SUMPRODUCT({m}*({xrng(val)}-{cmu}{r})^2)/({cn}{r}-1)),"")', fmt)

    # B. Carry 矩陣
    ws.cell(row=CM_HEAD - 2, column=1, value="B. Carry 分群（評等 × 天期）").font = BOLD
    for j, lab in ((2, "檔數"), (7, "OAS 平均"), (12, "OAS 標準差")):
        ws.cell(row=CM_HEAD - 1, column=j, value=lab).font = BOLD
    ws.cell(row=CM_HEAD, column=1, value="評等分群")
    for k, mb in enumerate(MAT_BUCKETS):
        for base in (2, 7, 12):
            ws.cell(row=CM_HEAD, column=base + k, value=mb)
    style_header(ws, CM_HEAD, 1, 16)
    for i, bkt in enumerate(BUCKETS):
        r = CM_R1 + i
        ws[f"A{r}"], ws[f"A{r}"].font = bkt, BLUE
        for k in range(len(MAT_BUCKETS)):
            cN, cM, cS = (get_column_letter(base + k) for base in (2, 7, 12))
            m = f"({xrng('valid')}=1)*({xrng('rb')}=$A{r})*({xrng('mb')}={cN}${CM_HEAD})"
            put(r, cN, f"=SUMPRODUCT({m})", "0")
            put(r, cM, f'=IF({cN}{r}>0,SUMPRODUCT({m}*{xrng("h_oas")})/{cN}{r},"")', "0.0")
            put(r, cS, f'=IF({cN}{r}>1,SQRT(SUMPRODUCT({m}*({xrng("h_oas")}-{cM}{r})^2)/({cN}{r}-1)),"")', "0.0")

    # C. 發行人曲線
    ws.cell(row=IS_HEAD - 1, column=1, value="C. 發行人曲線（對應「發行人」表 D 欄）").font = BOLD
    for j, h in enumerate(heads[:8], start=1):
        ws.cell(row=IS_HEAD, column=j, value="債券 Ticker" if j == 1 else h)
    style_header(ws, IS_HEAD, 1, 8)
    for i in range(IS_R2 - IS_R1 + 1):
        r, ir = IS_R1 + i, ISS_FIRST + i
        put(r, "A", f'=IF({I}!$D{ir}="","",{I}!$D{ir})', "General")
        fit_cols(r, f"({xrng('valid')}=1)*({xrng('ticker')}=$A{r})*($A{r}<>\"\")", MIN_ISS)
    for j, w in enumerate([14] + [11] * 18, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "B5"


def build_ranking(wb):
    ws = wb.create_sheet("排行")
    ws["A1"], ws["A1"].font = "綜合分數排行與分類概覽", TITLE
    ws["A2"] = "全部有效債券依綜合分數排序（不分評等）。要看單一評等分群，請在「債券」表用篩選。"
    ws["A2"].font = GREY
    fields = [("排名", None), ("債券代碼", "id"), ("債券簡稱", "des"), ("分類", "cat"), ("評等", "rtg"),
              ("剩餘年限", "yrs"), ("OAS (bp)", "oas"), ("DTS", "dts"), ("Carry z", "cz"), ("Value z", "vz"),
              ("Momentum z", "mz"), ("股票動能 z", "ez"), ("綜合分數", "score"), ("相對發行人\n曲線 (bp)", "ires")]
    fmts = {"yrs": "0.0", "oas": "0", "dts": "#,##0", "cz": "0.00", "vz": "0.00", "mz": "0.00", "ez": "0.00",
            "score": "0.00", "ires": "0.0;-0.0;0.0"}

    def block(top_row, title, func, helper):
        ws.cell(row=top_row, column=1, value=title).font = BOLD
        for j, (h, _) in enumerate(fields, start=1):
            ws.cell(row=top_row + 1, column=j, value=h)
        style_header(ws, top_row + 1, 1, len(fields))
        ws.row_dimensions[top_row + 1].height = 32
        for k in range(1, TOP_N + 1):
            r = top_row + 1 + k
            ws.cell(row=r, column=1, value=k).font = BLACK
            pos = f"MATCH({func}({xrng(helper)},$A{r}),{xrng(helper)},0)"
            guard = f"ABS({func}({xrng(helper)},$A{r}))>1E8"
            for j, (h, key) in enumerate(fields[1:], start=2):
                x = ws.cell(row=r, column=j,
                            value=f'=IF({guard},"",IF(INDEX({xrng(key)},{pos})="","",INDEX({xrng(key)},{pos})))')
                x.font = GREEN
                if key in fmts:
                    x.number_format = fmts[key]

    block(4, f"偏便宜／偏強勢 前 {TOP_N} 名（加碼研究清單）", "LARGE", "h_hi")
    block(6 + TOP_N + 1, f"偏貴／偏弱勢 前 {TOP_N} 名（減碼研究清單）", "SMALL", "h_lo")

    top = 6 + 2 * (TOP_N + 2) + 2
    ws.cell(row=top, column=1, value="分類 × 評等分群：平均 OAS (bp) 與檔數").font = BOLD
    cats = list(dict.fromkeys(i[0] for i in ISSUERS))
    buckets = ["AA以上", "A", "BBB", "HY"]
    heads = ["分類"] + [f"{b}\nOAS" for b in buckets] + [f"{b}\n檔數" for b in buckets] + ["全部\n加權 OAS", "全部\n檔數"]
    for j, h in enumerate(heads, start=1):
        ws.cell(row=top + 1, column=j, value=h)
    style_header(ws, top + 1, 1, len(heads))
    ws.row_dimensions[top + 1].height = 32
    for i, cat in enumerate(cats):
        r = top + 2 + i
        ws.cell(row=r, column=1, value=cat).font = BLUE
        for j, b in enumerate(buckets):
            x = ws.cell(row=r, column=2 + j,
                        value=f'=IFERROR(AVERAGEIFS({xrng("oas")},{xrng("cat")},$A{r},{xrng("rb")},"{b}",{xrng("valid")},1),"")')
            x.font, x.number_format = GREEN, "0"
            y = ws.cell(row=r, column=2 + len(buckets) + j,
                        value=f'=COUNTIFS({xrng("cat")},$A{r},{xrng("rb")},"{b}",{xrng("valid")},1)')
            y.font, y.number_format = GREEN, "0;-0;-"
        m = f"({xrng('cat')}=$A{r})*({xrng('valid')}=1)"
        x = ws.cell(row=r, column=2 + 2 * len(buckets),
                    value=f'=IF(SUMPRODUCT({m}*{xrng("h_amt")})>0,SUMPRODUCT({m}*{xrng("h_amt")}*{xrng("h_oas")})/SUMPRODUCT({m}*{xrng("h_amt")}),"")')
        x.font, x.number_format = GREEN, "0"
        y = ws.cell(row=r, column=3 + 2 * len(buckets), value=f"=SUMPRODUCT({m})")
        y.font, y.number_format = GREEN, "0;-0;-"
    widths = [8, 22, 24, 18, 9, 10, 10, 10, 10, 10, 11, 11, 10, 12]
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    for blk in (6, 6 + TOP_N + 3):
        ws.conditional_formatting.add(
            f"M{blk}:M{blk + TOP_N - 1}",
            ColorScaleRule(start_type="num", start_value=-1.5, start_color="F8696B",
                           mid_type="num", mid_value=0, mid_color="FFFFFF",
                           end_type="num", end_value=1.5, end_color="63BE7B"))


def build_search(wb, sample: bool):
    ws = wb.create_sheet("搜尋")
    ws["A1"], ws["A1"].font = "用 BSRCH 抓債券清單", TITLE
    lines = [
        "1. 在終端機輸入 SRCH <GO>，依「說明」表的條件建立債券搜尋並儲存，名稱填在 B6。",
        "2. Bloomberg 頁籤 → Refresh 後，A8 的 BSRCH 公式會列出搜尋結果。",
        "3. 把 A 欄結果複製 →「債券」表 A5 起「選擇性貼上 → 值」。搜尋結果順序可能變動，貼成值比較穩定。",
    ]
    for i, t in enumerate(lines, start=2):
        ws.cell(row=i, column=1, value=t).font = BLACK
    ws["A6"], ws["A6"].font = "SRCH 儲存名稱", BOLD
    ws["B6"] = "AI_USD"
    ws["B6"].font, ws["B6"].fill = BLUE, YELLOW
    ws["A8"] = "（模擬模式不產生 BSRCH）" if sample else '=BSRCH("FI:"&B6)'
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 14


def build_readme(wb):
    ws = wb.active
    ws.title = "說明"
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 120
    rows = [
        ("美元 AI 相關公司債監控（Bloomberg Excel Add-in）", TITLE),
        ("用途：追蹤 AI 主流發行人的美元債，計算 Carry、Value、Momentum、股票動能四個因子，產出加碼／減碼研究清單。方法見 repo 的 credit-quant.md。", None),
        ("本活頁簿只是研究工具，訊號不構成交易建議；實際交易請依公司授權與內控流程。", GREY),
        ("", None),
        ("■ 工作表", BOLD),
        ("說明：本頁。", None),
        ("參數：評價日、權重、門檻、評等對照表。", None),
        ("發行人：24 家 AI 相關公司（超大型雲端、半導體、伺服器／網通、軟體、資料中心、AI 雲端），含股票、財報、CDS 與債券彙總。", None),
        ("債券：債券清單與所有計算（每列一檔，最多 400 檔）。", None),
        ("排行：綜合分數前 20／後 20 名，以及「分類 × 評等」的平均 OAS。", None),
        ("分群統計：評等曲線、發行人曲線、各因子的平均與標準差（自動計算，債券表查表引用）。", None),
        ("搜尋：用 BSRCH 從 SRCH 儲存的搜尋抓債券清單。", None),
        ("", None),
        ("■ 使用步驟", BOLD),
        ("1. 開啟 Bloomberg 終端機並登入，Excel 的 Bloomberg 頁籤要能使用。", None),
        ("2. 建立債券清單（擇一）：", None),
        ("   A. SRCH <GO> 建立搜尋，建議條件：幣別 = USD；發行人 = 「發行人」表的公司（可用 Ultimate Parent 篩選，才會包含子公司發行的債）；", None),
        ("      票息型態 = Fixed；求償順位 = Senior Unsecured；排除可轉債、私募；未到期；流通在外 ≥ 5 億美元。存成 AI_USD，再用「搜尋」表的 BSRCH 取出。", None),
        ("   B. 個別公司：<股票代碼> DDIS <GO>（債務分布）匯出債券清單。", None),
        ("   把代碼貼到「債券」表 A 欄（A5 起）。格式可以是「US1234567890 Corp」或「/isin/US1234567890」；BSRCH 的結果可直接用。", None),
        ("3. 檢查黃底的彭博欄位代碼（「債券」與「發行人」表第 3 列）。欄位名稱可能因權限或版本不同，有錯誤請用 FLDS <GO> 搜尋正確代碼後修改。", None),
        ("4. 檢查「發行人」表 D 欄（債券 Ticker）：要和債券 TICKER 欄位一致才對得上。標「請確認」的（TSMC、SK hynix、CoreWeave）請在終端機用 DES 查發債主體的代碼。", None),
        ("   「債券」表的分類若出現「未對應」，代表該券的 TICKER 不在發行人表中，可以新增一列發行人。", None),
        ("5. Bloomberg 頁籤 → Refresh → Refresh Workbook。第一次會比較慢（BDH 取歷史 OAS）。", None),
        ("6. 看「排行」表與「發行人」表；細節到「債券」表篩選「訊號」欄。", None),
        ("", None),
        ("■ 顏色", BOLD),
        ("藍字黃底：可以輸入或修改的儲存格。黑字：公式。綠字：引用其他工作表的公式。灰字標題：輔助欄（給公式用，請勿刪除）。", None),
        ("", None),
        ("■ 因子定義", BOLD),
        ("Carry z：在「評等分群 × 天期分群」內，OAS 相對組平均的 z 分數。分群不足 5 檔時改用評等分群。", None),
        ("Value z：每個評等分群擬合 OAS = a + b·ln(存續期間)，計算「實際 OAS − 擬合 OAS」÷ 擬合 OAS（相對便宜 %），再在評等分群內做 z 分數。", None),
        ("Momentum z：過去 6 個月到 1 個月前的近似超額報酬（5 個月 carry − 存續期間 × OAS 變化），除以 DTS 後在評等分群內做 z 分數。", None),
        ("股票動能 z：發行公司股價 6 個月報酬，在評等分群內做 z 分數。投資等級債的動能以股票訊號較有效（Gebhardt 等, 2005）。", None),
        ("綜合分數：四個 z 分數依「參數」表權重加權平均；評等分群內排名前 20% 標「加碼候選」、後 20% 標「減碼候選」。", None),
        ("相對發行人曲線：同一發行人的債券擬合 OAS = a + b·ln(存續期間)，正值 = 相對自家曲線便宜，適合在同一發行人內換券。", None),
        ("DTS = 利差存續期間 × OAS，代表信用風險大小；發行人表的加權 DTS 可用來比較各公司的部位風險。", None),
        ("", None),
        ("■ 注意事項", BOLD),
        ("• 標的數少（約 20 多家發行人），因子分數主要反映相對位置，不宜當作嚴格的統計結論；建議每月記錄分數，累積後回頭驗證。", None),
        ("• 相同發行人的多檔債券分數高度相關，加碼時請注意單一發行人集中度。", None),
        ("• 利差存續期間以修正存續期間近似；可贖回（非 make-whole）債券應改用 OAS 存續期間。", None),
        ("• 財報欄位為最近一期年度或滾動資料，依 Bloomberg 預設；需要特定期間請加 BEST_FPERIOD_OVERRIDE 等覆寫參數。", None),
        ("• BDH 歷史資料會計入 Bloomberg 資料用量，債券多時建議月更新，不要頻繁全部重新整理。", None),
        ("• 「發行人」表的分類與備註是整理時的概略描述，請依最新資訊更新。", None),
        ("", None),
        ("■ 可以延伸的方向", BOLD),
        ("• 加入電力與公用事業（AI 資料中心用電）、AI 相關 SPV／資產擔保融資等發行人。", None),
        ("• 每月把「債券」表另存一份歷史快照，累積後用 repo 中的 credit_quant Python 範例做回測。", None),
    ]
    for i, (t, f) in enumerate(rows, start=1):
        x = ws.cell(row=i, column=2, value=t)
        x.font = f or BLACK
        x.alignment = Alignment(wrap_text=True, vertical="top")


def build(path: Path, sample: bool = False):
    wb = Workbook()
    build_readme(wb)
    build_params(wb)
    build_issuers(wb, sample)
    build_bonds(wb, sample)
    build_stats(wb)
    build_ranking(wb)
    build_search(wb, sample)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for x in row:
                if x.font and x.font.name != FONT:
                    x.font = Font(name=FONT, bold=x.font.bold, color=x.font.color, size=x.font.size)
    wb.calculation.fullCalcOnLoad = True  # 開檔時全部重算
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--sample", metavar="PATH", help="產生以模擬數字取代 Bloomberg 公式的測試版")
    args = ap.parse_args()
    if args.sample:
        print(build(Path(args.sample), sample=True))
    else:
        print(build(Path(__file__).with_name("AI_USD_Credit_Monitor.xlsx")))
